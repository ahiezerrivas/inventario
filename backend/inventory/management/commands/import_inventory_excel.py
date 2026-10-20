from __future__ import annotations

import unicodedata
from pathlib import Path
from typing import Iterable

from django.core.management.base import BaseCommand, CommandError
from openpyxl import load_workbook

from inventory.models import InventoryItem

SHEET_MAP = {
    "Inventario Global": "Global",
    "Inventario Sede": "Sede",
    "Inventario Oficina": "Oficina",
}

HEADER_ALIASES = {
    "asset_type": {"tipo de equipo", "tipo equipo"},
    "model": {"modelo"},
    "serial": {"serial equipo", "serial"},
    "assignee": {"usuario"},
    "office": {"oficina", "oficina unidad", "oficina / unidad", "unidad oficina", "unidad responsable"},
    "status": {"status", "estatus"},
    "asset_code": {
        "cod.",
        "cod",
        "cod. de oficina",
        "cod de oficina",
        "codigo de oficina",
        "cod. oficina",
        "cod. adm.",
        "cod. adm",
        "codigo de unidad",
        "cod. unidad",
    },
}

SOURCE_BY_LOCATION = {
    "sede": "Sede",
    "oficina": "Oficina",
    "global": "Global",
}


def normalize(text: object) -> str:
    if text is None:
        return ""
    raw = str(text).strip().lower()
    raw = unicodedata.normalize("NFKD", raw)
    raw = "".join(ch for ch in raw if not unicodedata.combining(ch))
    return " ".join(raw.replace("_", " ").replace("-", " ").replace("/", " ").split())


def first_header_row(rows: Iterable[tuple], max_scan: int = 40) -> int:
    best_row = 1
    best_score = -1
    for idx, row in enumerate(rows, start=1):
        if idx > max_scan:
            break
        score = 0
        for value in row:
            val = normalize(value)
            if val and not val.isdigit():
                score += 1
        if score > best_score:
            best_score = score
            best_row = idx
    return best_row


def extract_value(row_values: dict[str, object], aliases: set[str]) -> str:
    for key, value in row_values.items():
        if key in aliases:
            return "" if value is None else str(value)
    return ""


def infer_source_sheet(row_values: dict[str, object]) -> str:
    location = normalize(row_values.get("ubicacion", ""))
    return SOURCE_BY_LOCATION.get(location, "Global")


def normalize_serial(value: object) -> str:
    return str(value or "").strip().upper()


def build_existing_serials() -> set[str]:
    existing: set[str] = set()
    for item in InventoryItem.objects.values("serial", "extra"):
        primary = normalize_serial(item["serial"])
        if primary:
            existing.add(primary)
        extra = item["extra"] or {}
        for key in ("serial equipo", "serial"):
            value = normalize_serial(extra.get(key))
            if value:
                existing.add(value)
    return existing


class Command(BaseCommand):
    help = "Importa Inventario Global/Sede/Oficina hacia SQLite"

    def add_arguments(self, parser):
        parser.add_argument("--file", required=True, help="Ruta del archivo .xlsx")
        parser.add_argument(
            "--replace",
            action="store_true",
            help="Borra registros previos de Global/Sede/Oficina antes de importar",
        )
        parser.add_argument(
            "--sheet",
            default=None,
            help="Importa solo la hoja indicada (ej. 'Inventario Global'). Si se omite, usa 'Detalle' o Global/Sede/Oficina.",
        )
        parser.add_argument(
            "--skip-duplicate-serials",
            action="store_true",
            help="Omite filas cuyo serial ya exista en la base de datos (o repetido dentro del archivo).",
        )
        parser.add_argument(
            "--infer-source",
            action="store_true",
            help="Clasifica cada fila por su columna 'Ubicacion' (Sede/Oficina/Global) en lugar de forzar el nombre de la hoja.",
        )

    def handle(self, *args, **options):
        excel_path = Path(options["file"]).expanduser().resolve()
        if not excel_path.exists():
            raise CommandError(f"No existe el archivo: {excel_path}")

        if options["replace"]:
            InventoryItem.objects.filter(source_sheet__in=["Global", "Sede", "Oficina"]).delete()

        wb = load_workbook(excel_path, data_only=True, read_only=True)

        total = 0
        skipped_duplicates = 0
        per_sheet: dict[str, int] = {"Global": 0, "Sede": 0, "Oficina": 0}
        buffer: list[InventoryItem] = []

        skip_duplicates = options["skip_duplicate_serials"]
        existing_serials = build_existing_serials() if skip_duplicates else set()
        if skip_duplicates:
            self.stdout.write(f"Validacion de seriales activa. Seriales existentes: {len(existing_serials)}.")

        if options["sheet"]:
            sheet_name = options["sheet"]
            sources_to_import = [(sheet_name, SHEET_MAP.get(sheet_name))]
            self.stdout.write(f"Modo de importacion: hoja unica '{sheet_name}'.")
        elif "Detalle" in wb.sheetnames:
            sources_to_import = [("Detalle", None)]
            self.stdout.write("Modo de importacion: hoja 'Detalle' (clasificacion por Ubicacion).")
        else:
            sources_to_import = list(SHEET_MAP.items())
            self.stdout.write("Modo de importacion: hojas Inventario Global/Sede/Oficina.")

        for sheet_name, fixed_source_sheet in sources_to_import:
            if sheet_name not in wb.sheetnames:
                self.stdout.write(self.style.WARNING(f"No existe hoja: {sheet_name}"))
                continue

            ws = wb[sheet_name]
            row_iter = ws.iter_rows(min_row=1, max_row=min(ws.max_row, 40), values_only=True)
            header_row = first_header_row(row_iter)

            headers = [normalize(cell) for cell in next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))]

            for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
                if not any(value is not None and str(value).strip() != "" for value in row):
                    continue

                row_values = {headers[i]: row[i] for i in range(min(len(headers), len(row))) if headers[i]}
                if options["infer_source"]:
                    source_sheet = infer_source_sheet(row_values)
                else:
                    source_sheet = fixed_source_sheet or infer_source_sheet(row_values)

                serial_value = extract_value(row_values, HEADER_ALIASES["serial"])
                if skip_duplicates:
                    normalized = normalize_serial(serial_value)
                    if normalized and normalized in existing_serials:
                        skipped_duplicates += 1
                        continue
                    if normalized:
                        existing_serials.add(normalized)

                item = InventoryItem(
                    source_sheet=source_sheet,
                    asset_code=extract_value(row_values, HEADER_ALIASES["asset_code"]),
                    asset_type=extract_value(row_values, HEADER_ALIASES["asset_type"]),
                    model=extract_value(row_values, HEADER_ALIASES["model"]),
                    serial=serial_value,
                    assignee=extract_value(row_values, HEADER_ALIASES["assignee"]),
                    office=extract_value(row_values, HEADER_ALIASES["office"]),
                    status=extract_value(row_values, HEADER_ALIASES["status"]).strip(),
                    extra={k: ("" if v is None else str(v)) for k, v in row_values.items()},
                )
                buffer.append(item)
                total += 1
                per_sheet[source_sheet] += 1

                if len(buffer) >= 1000:
                    InventoryItem.objects.bulk_create(buffer, batch_size=1000)
                    buffer.clear()

        if buffer:
            InventoryItem.objects.bulk_create(buffer, batch_size=1000)

        self.stdout.write(self.style.SUCCESS("Importacion completada."))
        self.stdout.write(f"Total insertados: {total}")
        self.stdout.write(f"Global: {per_sheet['Global']}")
        self.stdout.write(f"Sede: {per_sheet['Sede']}")
        self.stdout.write(f"Oficina: {per_sheet['Oficina']}")
        if skip_duplicates:
            self.stdout.write(f"Omitidos por serial duplicado: {skipped_duplicates}")

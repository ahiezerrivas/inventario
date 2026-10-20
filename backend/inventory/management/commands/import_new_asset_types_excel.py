from __future__ import annotations

import unicodedata
from pathlib import Path
from typing import Iterable

from django.core.management.base import BaseCommand, CommandError
from openpyxl import load_workbook

from inventory.models import InventoryItem

SHEET_MAP = {
    "Inventario Global": "Global",
    "Inventario Sede": "Sede",
    "Inventario Oficina": "Oficina",
}

HEADER_ALIASES = {
    "asset_type": {"tipo de equipo", "tipo equipo"},
    "model": {"modelo"},
    "serial": {"serial"},
    "assignee": {"usuario"},
    "office": {"oficina", "oficina unidad", "oficina / unidad", "unidad responsable"},
    "asset_code": {
        "cod.",
        "cod",
        "cod. de oficina",
        "cod de oficina",
        "codigo de oficina",
        "cod. oficina",
        "cod. adm.",
        "cod. adm",
        "codigo de unidad",
        "cod. unidad",
    },
}

SOURCE_BY_LOCATION = {
    "sede": "Sede",
    "oficina": "Oficina",
    "global": "Global",
}


def normalize(text: object) -> str:
    if text is None:
        return ""
    raw = str(text).strip().lower()
    raw = unicodedata.normalize("NFKD", raw)
    raw = "".join(ch for ch in raw if not unicodedata.combining(ch))
    return " ".join(raw.replace("_", " ").replace("-", " ").replace("/", " ").split())


def canonical_asset_type(asset_type: str) -> str:
    return str(asset_type or "").strip().lower()


def first_header_row(rows: Iterable[tuple], max_scan: int = 40) -> int:
    best_row = 1
    best_score = -1
    for idx, row in enumerate(rows, start=1):
        if idx > max_scan:
            break
        score = 0
        for value in row:
            val = normalize(value)
            if val and not val.isdigit():
                score += 1
        if score > best_score:
            best_score = score
            best_row = idx
    return best_row


def extract_value(row_values: dict[str, object], aliases: set[str]) -> str:
    for key, value in row_values.items():
        if key in aliases:
            return "" if value is None else str(value).strip()
    return ""


def infer_source_sheet(row_values: dict[str, object]) -> str:
    location = normalize(row_values.get("ubicacion", ""))
    if not location:
        return ""
    if "oficina" in location:
        return "Oficina"
    if "sede" in location:
        return "Sede"
    return SOURCE_BY_LOCATION.get(location, "")


class Command(BaseCommand):
    help = "Importa SOLO filas con tipo de equipo nuevo y serial no existente desde un Excel"

    def handle(self, *args, **options):
        excel_path = Path(options["file"]).expanduser().resolve()
        apply_changes = bool(options["apply"])

        if not excel_path.exists():
            raise CommandError(f"No existe el archivo: {excel_path}")

        existing_types = {
            canonical_asset_type(t)
            for t in InventoryItem.objects.values_list("asset_type", flat=True)
            if t and str(t).strip()
        }
        existing_serials = {
            str(s).strip().lower()
            for s in InventoryItem.objects.values_list("serial", flat=True)
            if s and str(s).strip()
        }

        total_rows = 0
        candidates_new_type = 0
        skipped_existing_type = 0
        skipped_missing_serial = 0
        skipped_existing_serial = 0
        skipped_duplicate_serial_in_file = 0
        to_insert = 0
        new_types_in_file: set[str] = set()

        seen_serials_in_file: set[str] = set()
        buffer: list[InventoryItem] = []

        wb = load_workbook(excel_path, data_only=True, read_only=True)
        if "Detalle" in wb.sheetnames:
            sources_to_import = [("Detalle", None)]
        else:
            sources_to_import = list(SHEET_MAP.items())

        for sheet_name, fixed_source_sheet in sources_to_import:
            if sheet_name not in wb.sheetnames:
                self.stdout.write(self.style.WARNING(f"No existe hoja: {sheet_name}"))
                continue

            ws = wb[sheet_name]
            row_iter = ws.iter_rows(min_row=1, max_row=min(ws.max_row, 40), values_only=True)
            header_row = first_header_row(row_iter)
            headers = [
                normalize(cell)
                for cell in next(
                    ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True)
                )
            ]

            for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
                if not any(value is not None and str(value).strip() != "" for value in row):
                    continue

                total_rows += 1
                row_values = {
                    headers[i]: row[i]
                    for i in range(min(len(headers), len(row)))
                    if headers[i]
                }

                raw_asset_type = extract_value(row_values, HEADER_ALIASES["asset_type"])
                asset_type = canonical_asset_type(raw_asset_type)
                if not asset_type:
                    skipped_existing_type += 1
                    continue

                if asset_type in existing_types:
                    skipped_existing_type += 1
                    continue

                candidates_new_type += 1
                new_types_in_file.add(asset_type)

                serial = extract_value(row_values, HEADER_ALIASES["serial"])
                serial_key = serial.lower()
                if not serial_key:
                    skipped_missing_serial += 1
                    continue

                if serial_key in existing_serials:
                    skipped_existing_serial += 1
                    continue

                if serial_key in seen_serials_in_file:
                    skipped_duplicate_serial_in_file += 1
                    continue

                seen_serials_in_file.add(serial_key)
                inferred_source = infer_source_sheet(row_values)
                source_sheet = inferred_source or fixed_source_sheet or "Global"
                asset_code = extract_value(row_values, HEADER_ALIASES["asset_code"])
                model = extract_value(row_values, HEADER_ALIASES["model"])
                assignee = extract_value(row_values, HEADER_ALIASES["assignee"])
                office = extract_value(row_values, HEADER_ALIASES["office"])
                extra = {k: ("" if v is None else str(v).strip()) for k, v in row_values.items()}
                extra["tipo de equipo"] = asset_type

                item = InventoryItem(
                    source_sheet=source_sheet,
                    asset_code=asset_code,
                    asset_type=asset_type,
                    model=model,
                    serial=serial,
                    assignee=assignee,
                    office=office,
                    extra=extra,
                )

                buffer.append(item)
                to_insert += 1

        self.stdout.write(self.style.SUCCESS("Analisis completado."))
        self.stdout.write(f"Archivo: {excel_path}")
        self.stdout.write(f"Filas leidas: {total_rows}")
        self.stdout.write(f"Filas candidatas por tipo nuevo: {candidates_new_type}")
        self.stdout.write(f"Tipos nuevos detectados: {len(new_types_in_file)}")
        if new_types_in_file:
            self.stdout.write(
                "Tipos nuevos: " + ", ".join(sorted(new_types_in_file))
            )
        self.stdout.write(f"Saltadas por tipo existente/vacio: {skipped_existing_type}")
        self.stdout.write(f"Saltadas por serial vacio: {skipped_missing_serial}")
        self.stdout.write(f"Saltadas por serial ya existente en DB: {skipped_existing_serial}")
        self.stdout.write(f"Saltadas por serial duplicado dentro del Excel: {skipped_duplicate_serial_in_file}")
        self.stdout.write(f"Insertables finales: {to_insert}")

        if not apply_changes:
            self.stdout.write(
                self.style.WARNING(
                    "Modo dry-run: no se inserto nada. Usa --apply para aplicar los cambios."
                )
            )
            return

        if buffer:
            InventoryItem.objects.bulk_create(buffer, batch_size=1000)

        self.stdout.write(
            self.style.SUCCESS(
                f"Importacion aplicada. Insertados: {to_insert}"
            )
        )

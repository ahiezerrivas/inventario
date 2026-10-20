import openpyxl
import traceback

SOURCE = r"c:/Users/admin-c606415/Documents/inventario/Prueba de Inventario avance.xlsx"
OUTPUT = r"c:/Users/admin-c606415/Documents/inventario/Prueba de Inventario avance_auto.xlsx"

with open("setup_auto_sync.log", "w", encoding="utf-8") as log:
    try:
        log.write("Inicio\n")
        wb = openpyxl.load_workbook(SOURCE)
        log.write("Workbook cargado\n")

        configs = [
            ("Inventario Sede", 10, "Sede"),
            ("Inventario Oficina", 4, "Oficina"),
        ]

        for sheet_name, start_row, location_value in configs:
            ws = wb[sheet_name]
            log.write(f"Procesando: {sheet_name}\n")

            # Clear previous table values from the dynamic area.
            for r in range(start_row, ws.max_row + 1):
                for c in range(1, ws.max_column + 1):
                    ws.cell(row=r, column=c).value = None

            formula = (
                "=FILTER('Inventario Global'!A4:Z1048576,"
                f"'Inventario Global'!$B$4:$B$1048576=\"{location_value}\""
                ',"Sin datos")'
            )

            ws.cell(row=start_row, column=1).value = formula

        wb.save(OUTPUT)
        log.write(f"Archivo creado: {OUTPUT}\n")
        print(f"Archivo creado: {OUTPUT}")
    except Exception:
        log.write("ERROR\n")
        log.write(traceback.format_exc())
        raise

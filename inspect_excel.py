import openpyxl

path = r"c:/Users/admin-c606415/Documents/inventario/Prueba de Inventario avance.xlsx"
wb = openpyxl.load_workbook(path)
print("SHEETS:", wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    print(f"{name} -> rows={ws.max_row}, cols={ws.max_column}")
    headers = [ws.cell(row=1, column=i).value for i in range(1, min(ws.max_column, 12) + 1)]
    print(" headers:", headers)
